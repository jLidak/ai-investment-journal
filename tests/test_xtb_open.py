
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.ingestion.xtb_open import read_open_positions


def create_workbook(
    tmp_path,
    header_row=9,
    invalid_header=False,
    volume=3,
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Open Positions"

    sheet["B1"] = "DEMO_ACCOUNT"

    headers = [
        "Product", "Instrument/Position", "Ticker",
        "Category", "Type", "Volume",
    ]

    if invalid_header:
        headers[5] = "Wrong Volume"

    for column, header in enumerate(headers, start=1):
        sheet.cell(header_row, column, header)

    rows = [
        ["My Trades", "Demo", "DEMO.US", "STOCK", "", volume],
        ["My Trades", "POS_001", "DEMO.US", "", "BUY", 2],
        ["My Trades", "POS_002", "DEMO.US", "", "BUY", 1],
        ["My Trades", "Global ETF", "GLOBAL.DE", "ETF", "", 2],
        ["My Trades", "POS_003", "GLOBAL.DE", "", "BUY", 2],
    ]

    for row_number, row in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(row, start=1):
            sheet.cell(row_number, column, value)

    path = tmp_path / "open_positions.xlsx"
    workbook.save(path)
    workbook.close()

    return path


@pytest.mark.parametrize("header_row", [9, 11])
def test_reads_summary_positions(tmp_path, header_row):
    path = create_workbook(tmp_path, header_row=header_row)

    result = read_open_positions(path)

    assert result.reported_account_id == "DEMO_ACCOUNT"
    assert result.quantities == {
        "DEMO.US": Decimal("3"),
        "GLOBAL.DE": Decimal("2"),
    }


def test_rejects_invalid_header(tmp_path):
    path = create_workbook(tmp_path, invalid_header=True)

    with pytest.raises(ValueError, match="header not found"):
        read_open_positions(path)


def test_rejects_negative_volume(tmp_path):
    path = create_workbook(tmp_path, volume=-3)

    with pytest.raises(ValueError, match="Invalid volume"):
        read_open_positions(path)
