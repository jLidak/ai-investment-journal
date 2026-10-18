
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


file_path = Path("data/private/xtb_sample.xlsx")

workbook = load_workbook(
    filename=file_path,
    read_only=True,
    data_only=True,
)

print("Dostępne arkusze:", workbook.sheetnames)

sheet = workbook["Cash Operations"]
sheet.reset_dimensions()

headers = [
    cell.value
    for cell in sheet[5]
]

print("Nagłówki:", headers)

operation_types = Counter()

for row in sheet.iter_rows(
    min_row=6,
    values_only=True,
):
    operation_type = row[0]

    if operation_type is None or operation_type == "Total":
        continue

    operation_types[operation_type] += 1

print("Liczba operacji:", sum(operation_types.values()))
print("Typy operacji:", operation_types)

workbook.close()
