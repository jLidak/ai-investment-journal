
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HEADER = (
    "Product",
    "Instrument/Position",
    "Ticker",
    "Category",
    "Type",
    "Volume",
)


@dataclass(frozen=True)
class OpenPositionsSnapshot:
    reported_account_id: str
    quantities: dict[str, Decimal]


def read_open_positions(
    file_path: str | Path,
) -> OpenPositionsSnapshot:
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Open Positions"]
        sheet.reset_dimensions()

        account_id = sheet["B1"].value

        if account_id is None or not str(account_id).strip():
            raise ValueError("Missing account ID")

        # The header may appear on different rows.
        header_row = None

        for row_number, row in enumerate(
            sheet.iter_rows(max_col=6, values_only=True),
            start=1,
        ):
            if tuple(row) == EXPECTED_HEADER:
                header_row = row_number
                break

        if header_row is None:
            raise ValueError("Open Positions header not found")

        quantities = {}

        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=header_row + 1,
                max_col=6,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            product, instrument, ticker, category, kind, volume = row

            # Instrument summary row, not an individual lot.
            if category in ("STOCK", "ETF") and kind in (None, ""):
                if not ticker or volume is None:
                    raise ValueError(
                        f"Missing ticker or volume in row {row_number}"
                    )

                try:
                    quantity = Decimal(str(volume))
                except InvalidOperation as exc:
                    raise ValueError(
                        f"Invalid volume in row {row_number}"
                    ) from exc

                if not quantity.is_finite() or quantity < 0:
                    raise ValueError(
                        f"Invalid volume in row {row_number}"
                    )

                ticker = str(ticker)

                quantities[ticker] = (
                    quantities.get(ticker, Decimal("0")) + quantity
                )

            # Individual position rows are already included in summaries.
            elif kind in ("BUY", "SELL"):
                continue

            elif ticker not in (None, ""):
                raise ValueError(
                    f"Unexpected position row {row_number}"
                )

        return OpenPositionsSnapshot(
            reported_account_id=str(account_id),
            quantities=quantities,
        )

    finally:
        workbook.close()
