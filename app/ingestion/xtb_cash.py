
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HEADERS = (
    "Type", "Instrument", "Ticker", "Category",
    "Time", "Amount", "ID", "Comment",
    "Product", "Position ID",
)


@dataclass(frozen=True)
class CashOperation:
    account_id: str
    operation_type: str
    instrument: str | None
    ticker: str | None
    category: str | None
    time_utc: datetime
    amount: Decimal
    external_id: str
    comment: str
    product: str | None
    position_id: str | None


def read_cash_operations(file_path: str | Path) -> list[CashOperation]:
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Cash Operations"]

        # XTB może zapisywać błędne wymiary arkusza.
        sheet.reset_dimensions()

        if sheet["A3"].value != "Date from (UTC)":
            raise ValueError("Nieznana strefa czasowa eksportu")

        headers = tuple(cell.value for cell in sheet[5])

        if headers != EXPECTED_HEADERS:
            raise ValueError("Nieoczekiwany format kolumn XTB")

        account_id = str(sheet["B1"].value)

        operations = []

        for row_number, row in enumerate(
                sheet.iter_rows(
                    min_row=6,
                    max_col=len(EXPECTED_HEADERS),
                    values_only=True,
                ),
                start=6,
        ):
            if not row or row[0] in (None, "Total"):
                continue

            if not isinstance(row[4], datetime):
                raise ValueError(
                    f"Nieprawidłowa data w wierszu {row_number}"
                )

            if not row[6] or row[5] is None:
                raise ValueError(
                    f"Brak ID lub kwoty w wierszu {row_number}"
                )

            operation = CashOperation(
                account_id=account_id,
                operation_type=str(row[0]),
                instrument=str(row[1]) if row[1] else None,
                ticker=str(row[2]) if row[2] else None,
                category=str(row[3]) if row[3] else None,
                time_utc=row[4].replace(tzinfo=timezone.utc),
                amount=Decimal(str(row[5])),
                external_id=str(row[6]),
                comment=str(row[7]) if row[7] else "",
                product=str(row[8]) if row[8] else None,
                position_id=str(row[9]) if row[9] else None,
            )

            operations.append(operation)

        return operations

    finally:
        workbook.close()
